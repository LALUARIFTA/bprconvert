"""Run smoke checks against every currently available Flask conversion handler."""
import io
import json

import fitz
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

import app as backend
from app import app

client = app.test_client()


def pdf(text="BPRCONVERT smoke", pages=1):
    doc = fitz.open()
    for index in range(pages):
        page = doc.new_page()
        page.insert_text(fitz.Point(50, 100), f"{text} {index + 1}")
    payload = io.BytesIO(doc.tobytes())
    doc.close()
    payload.seek(0)
    return payload


def png(color=(255, 0, 0)):
    image = Image.new("RGB", (120, 80), color=color)
    payload = io.BytesIO()
    image.save(payload, "PNG")
    payload.seek(0)
    return payload


def docx():
    doc = Document()
    doc.add_paragraph("Word smoke")
    payload = io.BytesIO()
    doc.save(payload)
    payload.seek(0)
    return payload


def xlsx():
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "value"])
    sheet.append(["smoke", 1])
    payload = io.BytesIO()
    book.save(payload)
    payload.seek(0)
    return payload
def pptx():
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    payload = io.BytesIO()
    deck.save(payload)
    payload.seek(0)
    return payload


def protected_pdf(password="secret"):
    source = pdf("Unlock")
    with fitz.open(stream=source, filetype="pdf") as doc:
        payload = io.BytesIO(doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password))
    payload.seek(0)
    return payload




def post(tool_id, data):
    response = client.post(f"/api/convert/{tool_id}", data=data, content_type="multipart/form-data")
    assert response.status_code == 200, (tool_id, response.status_code, response.data[:200])
    assert response.data, tool_id
    return response


backend.ai_chat = lambda prompt, document_text, language=None: f"AI smoke: {language or 'summary'}: {document_text[:20]}"

assert client.get("/api/health").status_code == 200
catalog = client.get("/api/tools").get_json()
assert len(catalog) == 40, len(catalog)
assert all({"id", "title", "accept", "multiple", "fields", "available"} <= set(item) for item in catalog)

payloads = {
    "pdf-to-txt": lambda: {"files": (pdf("Text"), "text.pdf")},
    "word-to-txt": lambda: {"files": (docx(), "word.docx")},
    "excel-to-csv": lambda: {"files": (xlsx(), "sheet.xlsx")},
    "csv-to-json": lambda: {"files": (io.BytesIO(b"name,value\nsmoke,1\n"), "data.csv")},
    "json-to-csv": lambda: {"files": (io.BytesIO(json.dumps([{"name": "smoke", "value": 1}]).encode()), "data.json")},
    "png-to-jpg": lambda: {"files": (png(), "image.png")},
    "jpg-to-png": lambda: {"files": (png(), "image.jpg")},
    "md-to-json": lambda: {"files": (io.BytesIO(b"# Smoke\nBody"), "doc.md")},
    "txt-to-md": lambda: {"files": (io.BytesIO(b"Smoke text"), "doc.txt")},
    "pdf-tool-0": lambda: {"files": [(pdf("A"), "a.pdf"), (pdf("B"), "b.pdf")]},
    "pdf-tool-1": lambda: {"files": (pdf("Split", 2), "split.pdf"), "mode": "range", "start": "1", "end": "1"},
    "pdf-tool-2": lambda: {"files": (pdf("Compress"), "compress.pdf")},
    "pdf-tool-3": lambda: {"files": (pdf("Word"), "word.pdf")},
    "pdf-tool-4": lambda: {"files": (pdf("Slides"), "slides.pdf")},
    "pdf-tool-5": lambda: {"files": (pdf("Excel"), "excel.pdf")},
    "pdf-tool-6": lambda: {"files": (docx(), "word.docx")},
    "pdf-tool-7": lambda: {"files": (pptx(), "slides.pptx")},
    "pdf-tool-8": lambda: {"files": (xlsx(), "sheet.xlsx")},
    "pdf-tool-9": lambda: {"files": (pdf("Edit"), "edit.pdf"), "text": "added", "page": "1", "x": "40", "y": "40"},
    "pdf-tool-10": lambda: {"files": (pdf("Image"), "image.pdf"), "page": "1"},
    "pdf-tool-11": lambda: {"files": [(png(), "a.png"), (png((0, 255, 0)), "b.png")]},
    "pdf-tool-12": lambda: {"files": (pdf("Sign"), "sign.pdf"), "signature": (png(), "sig.png"), "page": "1"},
    "pdf-tool-13": lambda: {"files": (pdf("Watermark"), "wm.pdf"), "text": "SMOKE"},
    "pdf-tool-14": lambda: {"files": (pdf("Rotate"), "rotate.pdf"), "angle": "90"},
    "pdf-tool-15": lambda: {"html": "<h1>Smoke</h1><p>HTML</p>"},
    "pdf-tool-16": lambda: {"files": (protected_pdf(), "locked.pdf"), "password": "secret"},
    "pdf-tool-17": lambda: {"files": (pdf("Protect"), "protect.pdf"), "password": "secret"},
    "pdf-tool-18": lambda: {"files": (pdf("Organize", 2), "organize.pdf"), "order": "2,1"},
    "pdf-tool-19": lambda: {"files": (pdf("PDFA"), "pdfa.pdf")},
    "pdf-tool-20": lambda: {"files": (pdf("Repair"), "repair.pdf")},
    "pdf-tool-21": lambda: {"files": (pdf("Numbers", 2), "numbers.pdf"), "position": "bottom-center"},
    "pdf-tool-22": lambda: {"files": [(png(), "scan.png")]},
    "pdf-tool-23": lambda: {"files": (pdf("OCR"), "ocr.pdf")},
    "pdf-tool-24": lambda: {"files": (pdf("Left"), "left.pdf"), "file2": (pdf("Right"), "right.pdf")},
    "pdf-tool-25": lambda: {"files": (pdf("Secret term"), "redact.pdf"), "term": "Secret"},
    "pdf-tool-26": lambda: {"files": (pdf("Crop"), "crop.pdf"), "top": "1", "right": "1", "bottom": "1", "left": "1"},
    "pdf-tool-27": lambda: {"files": (pdf("Forms"), "forms.pdf")},
    "pdf-tool-28": lambda: {"text": "Summarize this smoke document."},
    "pdf-tool-29": lambda: {"text": "Translate this smoke document.", "language": "Indonesia"},
    "pdf-tool-30": lambda: {"files": (pdf("Markdown"), "markdown.pdf")},
}

for item in catalog:
    if not item["available"]:
        continue
    assert item["id"] in payloads, item
    post(item["id"], payloads[item["id"]]())

print(f"catalog: {len(catalog)} tools; exercised available handlers")
