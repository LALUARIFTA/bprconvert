"""BPRCONVERT engine.

Operasi PDF berjalan lokal dengan PyMuPDF/pypdf/reportlab.
Operasi AI meneruskan teks ke API OpenAI-compatible yang dikonfigurasi lewat env.
Kredensial hanya dibaca dari environment server dan tidak pernah dikirim ke browser.
"""

import io
import json
import math
import os
import shutil
import subprocess
import tempfile
import uuid
import zipfile

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

import fitz
import numpy as np
import requests
import xlsxwriter
from docx import Document
from flask import Flask, jsonify, request, send_file as flask_send_file
from flask_cors import CORS
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
OFFICE_BINARY = os.environ.get("LIBREOFFICE_BIN", "")
GS_BINARY = os.environ.get("GHOSTSCRIPT_BIN", "")
CONVERTER_TIMEOUT = int(os.environ.get("CONVERTER_TIMEOUT", "120"))

def _find_binary(configured: str, names: tuple[str, ...], fallbacks: tuple[str, ...] = ()) -> str | None:
    candidates = [configured] if configured else []
    candidates.extend(names)
    candidates.extend(fallbacks)
    for candidate in candidates:
        if not candidate:
            continue
        found = candidate if os.path.isabs(candidate) and os.path.isfile(candidate) else shutil.which(candidate)
        if found:
            return found
    return None

def office_binary() -> str | None:
    return _find_binary(OFFICE_BINARY, ("soffice", "libreoffice"), (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files\LibreOffice\program\soffice.com",
    ))

def ghostscript_binary() -> str | None:
    return _find_binary(GS_BINARY, ("gswin64c", "gs"), (
        r"C:\Program Files\gs\bin\gswin64c.exe",
        r"C:\Program Files\gs\gs10.08.0\bin\gswin64c.exe",
    ))

def _external_convert(binary: str, args: list[str], source: bytes, source_name: str, output_ext: str) -> bytes:
    with tempfile.TemporaryDirectory(prefix="bprconvert-") as temp:
        source_path = os.path.join(temp, source_name)
        with open(source_path, "wb") as handle:
            handle.write(source)
        subprocess.run([binary, *args], cwd=temp, check=True, timeout=CONVERTER_TIMEOUT,
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        outputs = [os.path.join(temp, name) for name in os.listdir(temp) if name.lower().endswith(output_ext)]
        if not outputs:
            raise RuntimeError("Konverter tidak menghasilkan file keluaran")
        with open(outputs[0], "rb") as handle:
            return handle.read()
from reportlab.pdfgen import canvas as rl_canvas

AI_API_BASE = os.environ.get("AI_API_BASE", "").rstrip("/")
AI_API_KEY = os.environ.get("AI_API_KEY", "")
AI_MODEL = os.environ.get("AI_MODEL", "my9")
MAX_CONTENT_LENGTH = int(os.environ.get("MAX_CONTENT_LENGTH", 100 * 1024 * 1024))

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_CONTENT_LENGTH
CORS(app, origins=os.environ.get("CORS_ORIGINS", "*"))

def send_file(file, name=None, **kwargs):
    """Return an attachment; supports the local legacy positional filename form."""
    if name is not None:
        kwargs.setdefault("download_name", name)
        kwargs.setdefault("as_attachment", True)
    return flask_send_file(file, **kwargs)


def out(name: str) -> str:
    return f"bprconvert_{uuid.uuid4().hex[:8]}_{name}"


def pdf_payload(doc: fitz.Document) -> io.BytesIO:
    buffer = io.BytesIO()
    buffer.write(doc.tobytes(garbage=4, deflate=True, clean=True))
    buffer.seek(0)
    return buffer


def open_pdf(stream: io.BytesIO, password: str | None = None) -> fitz.Document:
    doc = fitz.open(stream=stream, filetype="pdf")
    if doc.needs_pass and not doc.authenticate(password or request.form.get("password", "")):
        doc.close()
        raise ValueError("Password PDF salah atau belum diisi")
    return doc


def fail(message: str, status: int = 400):
    return jsonify({"error": message}), status


def document_file(storage, field: str = "files") -> io.BytesIO | None:
    item = storage.files.get(field) if storage else None
    if not item or not item.filename:
        return None
    return io.BytesIO(item.read())


def ai_chat(prompt: str, document_text: str, language: str | None = None) -> str:
    if not AI_API_KEY or not AI_API_BASE:
        raise RuntimeError("AI_API_KEY / AI_API_BASE belum dikonfigurasi di server")
    final_prompt = f"{prompt} Gunakan bahasa keluaran: {language}." if language else prompt
    body = {"model": AI_MODEL, "messages": [{"role": "system", "content": "Kamu mesin pemroses dokumen BPRCONVERT. Jawab ringkas dan akurat."}, {"role": "user", "content": f"{final_prompt}\n\nIsi dokumen:\n{document_text}"}]}
    response = requests.post(f"{AI_API_BASE}/chat/completions", headers={"Authorization": f"Bearer {AI_API_KEY}", "Content-Type": "application/json"}, json=body, timeout=180)
    response.raise_for_status()
    if "text/event-stream" in response.headers.get("content-type", ""):
        chunks = []
        for line in response.text.splitlines():
            if line.startswith("data:") and line[5:].strip() != "[DONE]":
                try:
                    chunks.append(json.loads(line[5:]).get("choices", [{}])[0].get("delta", {}).get("content", ""))
                except json.JSONDecodeError:
                    continue
        return "".join(chunks).strip()
    data = response.json()
    return data["choices"][0]["message"]["content"].strip()


def extract_text(stream: io.BytesIO, password: str | None = None) -> str:
    with open_pdf(stream, password) as doc:
        return "\n".join(page.get_text() for page in doc)


def send_pdf(doc: fitz.Document, name: str):
    return send_file(pdf_payload(doc), out(name), mimetype="application/pdf")


def zip_response(entries):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries:
            archive.writestr(name, payload)
    buffer.seek(0)
    return send_file(buffer, out("split.zip"), mimetype="application/zip")


def images_to_pdf(storages, name: str):
    stream = io.BytesIO()
    pdf = rl_canvas.Canvas(stream, pagesize=A4)
    width, height = A4
    for item in storages:
        image = Image.open(io.BytesIO(item.read())).convert("RGB")
        ratio = min(width / image.width, height / image.height, 1)
        size = (image.width * ratio, image.height * ratio)
        pdf.drawImage(ImageReader(image), (width - size[0]) / 2, (height - size[1]) / 2, *size)
        pdf.showPage()
    pdf.save()
    stream.seek(0)
    return send_file(stream, out(name), mimetype="application/pdf")


@app.get("/api/health")
def health():
    return jsonify({"status": "ok", "ai_configured": bool(AI_API_KEY and AI_API_BASE), "model": AI_MODEL})


@app.post("/api/merge-pdf")
def merge_pdf():
    files = request.files.getlist("files")
    if len(files) < 2:
        return fail("Pilih minimal dua file PDF untuk digabung")
    merged = fitz.open()
    for item in files:
        merged.insert_pdf(open_pdf(io.BytesIO(item.read())))
    return send_pdf(merged, "merged.pdf")


@app.post("/api/split-pdf")
def split_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    mode = request.form.get("mode", "range")
    with open_pdf(stream) as doc:
        if mode == "per-page":
            entries = []
            for index, page in enumerate(doc):
                single = fitz.open()
                single.insert_pdf(doc, from_page=index, to_page=index)
                entries.append((f"page-{index + 1}.pdf", pdf_payload(single).getvalue()))
            return zip_response(entries)
        start = max(int(request.form.get("start", 1)), 1)
        end = min(int(request.form.get("end", doc.page_count)), doc.page_count)
        if start > end:
            return fail("Halaman awal lebih besar dari halaman akhir")
        range_doc = fitz.open()
        range_doc.insert_pdf(doc, from_page=start - 1, to_page=end - 1)
        return send_pdf(range_doc, f"split-{start}-{end}.pdf")


@app.post("/api/compress-pdf")
def compress_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    mode = request.form.get("mode", "light")
    with open_pdf(stream) as doc:
        if mode == "extreme":
            compressed = fitz.open()
            for page in doc:
                pix = page.get_pixmap(dpi=96)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                image_buffer = io.BytesIO()
                img.save(image_buffer, format="JPEG", quality=55)
                new_page = compressed.new_page(width=page.rect.width, height=page.rect.height)
                new_page.insert_image(new_page.rect, stream=image_buffer.getvalue())
            return send_pdf(compressed, "compressed.pdf")
        doc.subset_fonts()
        return send_pdf(doc, "compressed.pdf")


@app.post("/api/pdf-to-text")
def pdf_to_text():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    text = extract_text(stream)
    return send_file(io.BytesIO(text.encode("utf-8")), out("extracted.txt"), mimetype="text/plain; charset=utf-8")


@app.post("/api/pdf-to-word")
def pdf_to_word():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    text = extract_text(stream)
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return send_file(buffer, out("document.docx"), mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.post("/api/pdf-to-powerpoint")
def pdf_to_powerpoint():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    presentation = Presentation()
    slide_width, slide_height = presentation.slide_width, presentation.slide_height
    with open_pdf(stream) as doc:
        for page in doc:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            lines = [line.strip() for line in page.get_text().splitlines() if line.strip()]
            title = lines[0][:120] if lines else f"Halaman {page.number + 1}"
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), slide_width - Inches(1), Inches(1))
            title_box.text_frame.text = title
            body_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), slide_width - Inches(1), slide_height - Inches(2))
            body_box.text_frame.text = "\n".join(lines[1:80])
            pix = page.get_pixmap(dpi=90)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            image_buffer = io.BytesIO()
            img.save(image_buffer, format="PNG")
            image_buffer.seek(0)
            slide.shapes.add_picture(image_buffer, 0, 0, width=slide_width, height=slide_height)
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    return send_file(buffer, out("slides.pptx"), mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")


@app.post("/api/pdf-to-excel")
def pdf_to_excel():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    buffer = io.BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet("Extracted")
    sheet.write(0, 0, "Halaman")
    sheet.write(0, 1, "Baris")
    sheet.write(0, 2, "Teks")
    row = 1
    with open_pdf(stream) as doc:
        for page in doc:
            for line_index, line in enumerate(page.get_text().splitlines(), start=1):
                if not line.strip():
                    continue
                sheet.write(row, 0, page.number + 1)
                sheet.write(row, 1, line_index)
                sheet.write(row, 2, line)
                row += 1
    book.close()
    buffer.seek(0)
    return send_file(buffer, out("data.xlsx"), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")








@app.post("/api/jpg-to-pdf")
@app.post("/api/scan-to-pdf")
def jpg_to_pdf():
    files = request.files.getlist("files")
    if not files:
        return fail("Tidak ada gambar")
    return images_to_pdf(files, "images.pdf")


@app.post("/api/pdf-to-jpg")
def pdf_to_jpg():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    page_index = max(int(request.form.get("page", 1)), 1)
    with open_pdf(stream) as doc:
        if page_index > doc.page_count:
            return fail("Nomor halaman melebihi jumlah halaman")
        pix = doc[page_index - 1].get_pixmap(dpi=int(request.form.get("dpi", 144)))
    return send_file(io.BytesIO(pix.tobytes("jpeg")), out(f"page-{page_index}.jpg"), mimetype="image/jpeg")


@app.post("/api/sign-pdf")
def sign_pdf():
    stream = document_file(request)
    signature = document_file(request, "signature")
    if not stream or not signature:
        return fail("File PDF dan gambar tanda tangan diperlukan")
    page_index = max(int(request.form.get("page", 1)), 1)
    x = float(request.form.get("x", 60))
    y = float(request.form.get("y", 120))
    width = float(request.form.get("width", 180))
    signature_img = Image.open(io.BytesIO(signature.read())).convert("RGBA")
    image_buffer = io.BytesIO()
    signature_img.save(image_buffer, format="PNG")
    with open_pdf(stream) as doc:
        if page_index > doc.page_count:
            return fail("Nomor halaman melebihi jumlah halaman")
        page = doc[page_index - 1]
        ratio = signature_img.height / signature_img.width
        rect = fitz.Rect(x, y, x + width, y + width * ratio)
        page.insert_image(rect, stream=image_buffer.getvalue())
        return send_pdf(doc, "signed.pdf")


@app.post("/api/watermark")
def watermark():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    text = request.form.get("text", "BPRCONVERT")
    opacity = float(request.form.get("opacity", 0.35))
    angle = math.radians(float(request.form.get("angle", 35)))
    matrix = fitz.Matrix(float(request.form.get("angle", 35)))
    with open_pdf(stream) as doc:
        for page in doc:
            point = fitz.Point(40, page.rect.height - 60)
            page.insert_text(
                point,
                text,
                fontsize=28,
                color=(0.9, 0.2, 0.18),
                morph=(point, matrix),
                fill_opacity=opacity,
                overlay=True,
            )
        return send_pdf(doc, "watermarked.pdf")


@app.post("/api/page-numbers")
def page_numbers():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    position = request.form.get("position", "bottom-center")
    with open_pdf(stream) as doc:
        total = doc.page_count
        for index, page in enumerate(doc, start=1):
            label = f"{index} / {total}"
            rect = page.rect
            width = page.insert_textbox(
                fitz.Rect(0, 10 if position == "top-center" else rect.height - 30, rect.width, 35 if position == "top-center" else rect.height - 5),
                label,
                fontsize=10,
                color=(0.2, 0.2, 0.2),
                align=fitz.TEXT_ALIGN_CENTER,
                rotate=0,
            )
            if width < 0:
                continue
        return send_pdf(doc, "numbered.pdf")


@app.post("/api/organize-pdf")
def organize_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    raw_order = [value.strip() for value in request.form.get("order", "").split(",") if value.strip()]
    with open_pdf(stream) as doc:
        total = doc.page_count
        try:
            order = [int(value) - 1 for value in raw_order]
        except ValueError:
            return fail("Format urutan harus angka dipisah koma, contoh: 3,1,2")
        invalid = [value for value in order if value < 0 or value >= total]
        if invalid or len(set(order)) != len(order):
            return fail(f"Urutan tidak valid untuk PDF dengan {total} halaman")
        organized = fitz.open()
        for index in order:
            organized.insert_pdf(doc, from_page=index, to_page=index)
        return send_pdf(organized, "organized.pdf")


@app.post("/api/repair-pdf")
def repair_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    try:
        with open_pdf(stream) as doc:
            repaired = fitz.open()
            for page in doc:
                repaired.insert_pdf(doc, from_page=page.number, to_page=page.number)
            return send_pdf(repaired, "repaired.pdf")
    except Exception as exc:  # noqa: BLE001
        return fail(f"PDF tidak dapat diperbaiki: {exc}", 422)


@app.post("/api/crop-pdf")
def crop_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    top = float(request.form.get("top", 0))
    right = float(request.form.get("right", 0))
    bottom = float(request.form.get("bottom", 0))
    left = float(request.form.get("left", 0))
    with open_pdf(stream) as doc:
        for page in doc:
            rect = page.rect
            new_rect = fitz.Rect(rect.x0 + left, rect.y0 + top, rect.x1 - right, rect.y1 - bottom)
            if new_rect.width <= 1 or new_rect.height <= 1 or not rect.contains(new_rect):
                return fail("Nilai crop terlalu besar atau keluar dari batas halaman")
            page.set_cropbox(new_rect)
        return send_pdf(doc, "cropped.pdf")


@app.post("/api/redact-pdf")
def redact_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    needle = request.form.get("term", "").strip()
    if not needle:
        return fail("Kata yang disensor wajib diisi")
    with open_pdf(stream) as doc:
        for page in doc:
            for rect in page.search_for(needle):
                page.add_redact_annot(rect, fill=(0, 0, 0))
            page.apply_redactions()
        return send_pdf(doc, "redacted.pdf")


@app.post("/api/compare-pdf")
def compare_pdf():
    first = document_file(request, "files")
    second = document_file(request, "file2")
    if not first or not second:
        return fail("Dua file PDF diperlukan")
    left = [line.strip() for line in extract_text(first).splitlines() if line.strip()]
    right = [line.strip() for line in extract_text(second).splitlines() if line.strip()]
    only_first = [line for line in left if line not in right]
    only_second = [line for line in right if line not in left]
    report = [
        "Hanya di file pertama:",
        *(f"- {line}" for line in only_first),
        "",
        "Hanya di file kedua:",
        *(f"- {line}" for line in only_second),
    ]
    return send_file(io.BytesIO("\n".join(report).encode("utf-8")), out("comparison.txt"), mimetype="text/plain; charset=utf-8")


@app.post("/api/html-to-pdf")
def html_to_pdf():
    html = request.form.get("html", "").strip()
    if not html:
        return fail("HTML wajib diisi")
    buffer = io.BytesIO()
    writer = fitz.DocumentWriter(buffer)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)
    story = fitz.Story(html=html)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    buffer.seek(0)
    return send_file(buffer, out("webpage.pdf"), mimetype="application/pdf")


@app.post("/api/pdf-to-markdown")
def pdf_to_markdown():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    blocks = []
    with open_pdf(stream) as doc:
        for page in doc:
            blocks.append(f"## Halaman {page.number + 1}")
            blocks.append(page.get_text())
    return send_file(io.BytesIO("\n\n".join(blocks).encode("utf-8")), out("document.md"), mimetype="text/markdown; charset=utf-8")


@app.post("/api/ai-summarize")
def ai_summarize():
    stream = document_file(request)
    text = request.form.get("text", "")
    if stream:
        text = extract_text(stream)
    if not text.strip():
        return fail("Kirim file PDF atau teks untuk diringkas")
    try:
        summary = ai_chat("Ringkas dokumen berikut dalam poin-poin singkat.", text)
    except Exception as exc:  # noqa: BLE001
        return fail(f"Gagal memanggil AI: {exc}", 502)
    return jsonify({"result": summary})


@app.post("/api/ai-translate")
def ai_translate():
    stream = document_file(request)
    text = request.form.get("text", "")
    if stream:
        text = extract_text(stream)
    if not text.strip():
        return fail("Kirim file PDF atau teks untuk diterjemahkan")
    language = request.form.get("language", "Inggris")
    try:
        translated = ai_chat("Terjemahkan seluruh dokumen berikut dan pertahankan struktur aslinya.", text, language=f"Bahasa {language}")
    except Exception as exc:  # noqa: BLE001
        return fail(f"Gagal memanggil AI: {exc}", 502)
    return jsonify({"result": translated})


@app.post("/api/ocr-pdf")
def ocr_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        return fail("Mesin OCR belum terpasang di server", 501)
    engine = RapidOCR()
    with open_pdf(stream) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=200)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            result, _ = engine(np.asarray(img))
            for item in result or []:
                box, text = item[0], item[1]
                xs = [point[0] for point in box]
                ys = [point[1] for point in box]
                scale = 72.0 / 200.0
                rect = fitz.Rect(min(xs) * scale, min(ys) * scale, max(xs) * scale, max(ys) * scale)
                page.insert_textbox(rect, text, fontsize=max(rect.height * 0.7, 6), render_mode=3)
        return send_pdf(doc, "searchable.pdf")


@app.post("/api/edit-pdf")
def edit_pdf():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    text = request.form.get("text", "").strip()
    highlight = request.form.get("highlight", "").strip()
    page_index = max(int(request.form.get("page", 1)), 1)
    x = float(request.form.get("x", 50))
    y = float(request.form.get("y", 80))
    size = float(request.form.get("size", 12))
    color_name = request.form.get("color", "black")
    colors = {"black": (0, 0, 0), "red": (0.9, 0.1, 0.1), "blue": (0.1, 0.3, 0.8), "green": (0.1, 0.6, 0.2)}
    with open_pdf(stream) as doc:
        if page_index > doc.page_count:
            return fail("Nomor halaman melebihi jumlah halaman")
        page = doc[page_index - 1]
        if text:
            page.insert_text(fitz.Point(x, y), text, fontsize=size, color=colors.get(color_name, (0, 0, 0)))
        if highlight:
            for rect in page.search_for(highlight):
                page.add_highlight_annot(rect)
        return send_pdf(doc, "edited.pdf")


@app.post("/api/pdf-forms")
def pdf_forms():
    stream = document_file(request)
    if not stream:
        return fail("Tidak ada file PDF")
    with open_pdf(stream) as doc:
        page = doc[0]
        label_y = 60
        for index, label in enumerate(["Nama", "Email", "Catatan"]):
            page.insert_text(fitz.Point(50, label_y), f"{label}:", fontsize=11)
            widget = fitz.Widget()
            widget.field_name = f"field_{index}"
            widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
            widget.rect = fitz.Rect(130, label_y - 12, 420, label_y + 2)
            page.add_widget(widget)
            label_y += 40
        checkbox = fitz.Widget()
        checkbox.field_name = "consent"
        checkbox.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX
        checkbox.rect = fitz.Rect(130, label_y - 12, 145, label_y + 2)
        page.add_widget(checkbox)
        page.insert_text(fitz.Point(150, label_y), "Saya menyetujui data di atas", fontsize=10)
        return send_pdf(doc, "forms.pdf")

def word_to_text():
    stream = document_file(request)
    if not stream: return fail("Tidak ada file Word")
    text = "\n".join(p.text for p in Document(stream).paragraphs)
    return send_file(io.BytesIO(text.encode()), download_name=out("document.txt"), as_attachment=True, mimetype="text/plain")

def excel_to_csv():
    import csv
    stream = document_file(request)
    if not stream: return fail("Tidak ada file Excel")
    output = io.StringIO(); writer = csv.writer(output)
    for row in load_workbook(stream, read_only=True, data_only=True).active.iter_rows(values_only=True): writer.writerow(row)
    return send_file(io.BytesIO(output.getvalue().encode()), download_name=out("sheet.csv"), as_attachment=True, mimetype="text/csv")

def csv_to_json():
    import csv
    stream = document_file(request)
    if not stream: return fail("Tidak ada file CSV")
    rows = list(csv.DictReader(io.StringIO(stream.getvalue().decode("utf-8-sig"))))
    return send_file(io.BytesIO(json.dumps(rows, ensure_ascii=False, indent=2).encode()), download_name=out("data.json"), as_attachment=True, mimetype="application/json")

def json_to_csv():
    import csv
    stream = document_file(request)
    if not stream: return fail("Tidak ada file JSON")
    data = json.loads(stream.getvalue()); rows = data if isinstance(data, list) else [data]
    fields = sorted({k for row in rows for k in row}); output = io.StringIO(); writer = csv.DictWriter(output, fieldnames=fields); writer.writeheader(); writer.writerows(rows)
    return send_file(io.BytesIO(output.getvalue().encode()), download_name=out("data.csv"), as_attachment=True, mimetype="text/csv")

def png_to_jpg():
    stream = document_file(request)
    if not stream: return fail("Tidak ada file PNG")
    img = Image.open(stream).convert("RGB"); buf = io.BytesIO(); img.save(buf, "JPEG", quality=90); buf.seek(0)
    return send_file(buf, download_name=out("image.jpg"), as_attachment=True, mimetype="image/jpeg")

def jpg_to_png():
    stream = document_file(request)
    if not stream: return fail("Tidak ada file JPG")
    img = Image.open(stream); buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
    return send_file(buf, download_name=out("image.png"), as_attachment=True, mimetype="image/png")

def md_to_json():
    stream = document_file(request)
    if not stream: return fail("Tidak ada file Markdown")
    text = stream.getvalue().decode(); data = {"lines": text.splitlines(), "text": text}
    return send_file(io.BytesIO(json.dumps(data, ensure_ascii=False, indent=2).encode()), download_name=out("document.json"), as_attachment=True, mimetype="application/json")

def txt_to_md():
    stream = document_file(request)
    if not stream: return fail("Tidak ada file teks")
    return send_file(io.BytesIO(("# Document\n\n" + stream.getvalue().decode()).encode()), download_name=out("document.md"), as_attachment=True, mimetype="text/markdown")

def rotate_pdf():
    stream = document_file(request)
    if not stream: return fail("Tidak ada file PDF")
    angle = int(request.form.get("angle", 90)) % 360
    with open_pdf(stream) as doc:
        for page in doc: page.set_rotation((page.rotation + angle) % 360)
        return send_pdf(doc, "rotated.pdf")

def protect_pdf():
    stream = document_file(request); password = request.form.get("password", "")
    if not stream or not password: return fail("File PDF dan password diperlukan")
    with open_pdf(stream) as doc:
        buf = io.BytesIO(doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=password, user_pw=password, permissions=0)); buf.seek(0)
        return send_file(buf, download_name=out("protected.pdf"), as_attachment=True, mimetype="application/pdf")

def unlock_pdf():
    stream = document_file(request); password = request.form.get("password", "")
    if not stream: return fail("Tidak ada file PDF")
    with open_pdf(stream, password) as doc: return send_pdf(doc, "unlocked.pdf")


def _office_to_pdf(extension: str):
    binary = office_binary()
    stream = document_file(request)
    if not binary or not stream:
        return fail("LibreOffice belum tersedia atau file tidak ada", 503)
    item = request.files.get("files")
    name = item.filename if item and item.filename else f"input{extension}"
    try:
        payload = _external_convert(binary, ["--headless", "--convert-to", "pdf", "--outdir", ".", name], stream.getvalue(), name, ".pdf")
    except (OSError, subprocess.SubprocessError, RuntimeError) as exc:
        app.logger.warning("Office conversion failed: %s", exc)
        return fail("Konversi Office gagal", 422)
    return send_file(io.BytesIO(payload), out("document.pdf"), mimetype="application/pdf")

def word_to_pdf(): return _office_to_pdf(".docx")
def powerpoint_to_pdf(): return _office_to_pdf(".pptx")
def excel_to_pdf(): return _office_to_pdf(".xlsx")

def pdf_to_pdfa():
    binary = ghostscript_binary()
    stream = document_file(request)
    if not binary or not stream:
        return fail("Ghostscript belum tersedia atau file tidak ada", 503)
    with tempfile.TemporaryDirectory(prefix="bprconvert-pdfa-") as temp:
        source, target = os.path.join(temp, "input.pdf"), os.path.join(temp, "output.pdf")
        with open(source, "wb") as handle: handle.write(stream.getvalue())
        args = ["-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite", "-dPDFA=2", "-dNOOUTERSAVE", "-sColorConversionStrategy=UseDeviceIndependentColor", "-sOutputFile=" + target, source]
        try:
            subprocess.run([binary, *args], check=True, timeout=CONVERTER_TIMEOUT, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            with open(target, "rb") as handle: payload = handle.read()
        except (OSError, subprocess.SubprocessError, FileNotFoundError):
            return fail("Konversi PDF/A gagal", 422)
    return send_file(io.BytesIO(payload), out("document-pdfa.pdf"), mimetype="application/pdf")

TOOL_DEFS = [
    ("pdf-to-txt", "PDF to Text", ".pdf", False, "pdf_to_text"),
    ("word-to-txt", "Word to Text", ".doc,.docx", False, "word_to_text"),
    ("excel-to-csv", "Excel to CSV", ".xls,.xlsx", False, "excel_to_csv"),
    ("csv-to-json", "CSV to JSON", ".csv", False, "csv_to_json"),
    ("json-to-csv", "JSON to CSV", ".json", False, "json_to_csv"),
    ("png-to-jpg", "PNG to JPG", ".png", False, "png_to_jpg"),
    ("jpg-to-png", "JPG to PNG", ".jpg,.jpeg", False, "jpg_to_png"),
    ("md-to-json", "Markdown to JSON", ".md", False, "md_to_json"),
    ("txt-to-md", "Text to Markdown", ".txt", False, "txt_to_md"),
]
_PDF_TOOL_NAMES = ["Merge PDF", "Split PDF", "Compress PDF", "PDF to Word", "PDF to PowerPoint", "PDF to Excel", "Word to PDF", "PowerPoint to PDF", "Excel to PDF", "Edit PDF", "PDF to JPG", "JPG to PDF", "Sign PDF", "Watermark", "Rotate PDF", "HTML to PDF", "Unlock PDF", "Protect PDF", "Organize PDF", "PDF to PDF/A", "Repair PDF", "Page numbers", "Scan to PDF", "OCR PDF", "Compare PDF", "Redact PDF", "Crop PDF", "PDF Forms", "AI Summarizer", "Translate PDF", "PDF to Markdown"]
_PDF_HANDLERS = ["merge_pdf", "split_pdf", "compress_pdf", "pdf_to_word", "pdf_to_powerpoint", "pdf_to_excel", "word_to_pdf", "powerpoint_to_pdf", "excel_to_pdf", "edit_pdf", "pdf_to_jpg", "jpg_to_pdf", "sign_pdf", "watermark", "rotate_pdf", "html_to_pdf", "unlock_pdf", "protect_pdf", "organize_pdf", "pdf_to_pdfa", "repair_pdf", "page_numbers", "jpg_to_pdf", "ocr_pdf", "compare_pdf", "redact_pdf", "crop_pdf", "pdf_forms", "ai_summarize", "ai_translate", "pdf_to_markdown"]
TOOL_DEFS.extend((f"pdf-tool-{i}", name, ".pdf", name in ("Merge PDF", "JPG to PDF", "Scan to PDF"), handler) for i, (name, handler) in enumerate(zip(_PDF_TOOL_NAMES, _PDF_HANDLERS)))

TOOL_FIELDS = {
    "split_pdf": [{"name": "mode", "label": "Mode", "type": "select", "default": "range", "options": [{"value": "range", "label": "Page range"}, {"value": "per-page", "label": "Every page"}]}, {"name": "start", "label": "Start page", "type": "number", "default": 1}, {"name": "end", "label": "End page", "type": "number"}],
    "pdf_to_jpg": [{"name": "page", "label": "Page number", "type": "number", "default": 1}],
    "jpg_to_pdf": [],
    "sign_pdf": [{"name": "signature", "label": "Signature image", "type": "file", "accept": ".png,.jpg,.jpeg", "required": True}],
    "watermark": [{"name": "text", "label": "Watermark text", "type": "text", "default": "BPRCONVERT", "required": True}],
    "rotate_pdf": [{"name": "angle", "label": "Rotation angle", "type": "select", "default": "90", "options": [{"value": "90", "label": "90°"}, {"value": "180", "label": "180°"}, {"value": "270", "label": "270°"}]}],
    "unlock_pdf": [{"name": "password", "label": "PDF password", "type": "password", "required": True}],
    "protect_pdf": [{"name": "password", "label": "Password", "type": "password", "required": True}],
    "organize_pdf": [{"name": "order", "label": "Page order (comma-separated)", "type": "text", "required": True}],
    "compare_pdf": [{"name": "file2", "label": "Second PDF", "type": "file", "accept": ".pdf", "required": True}],
    "redact_pdf": [{"name": "term", "label": "Text to redact", "type": "text", "required": True}],
    "html_to_pdf": [{"name": "html", "label": "HTML content", "type": "textarea", "required": True}],
    "ai_summarize": [{"name": "text", "label": "Text (when not uploading PDF)", "type": "textarea"}],
    "ai_translate": [{"name": "language", "label": "Target language", "type": "text", "default": "Inggris"}],
    "pdf_to_pdfa": [],
}
TOOL_DESCRIPTIONS = {"pdf_to_word": "Extracts text into DOCX; original layout is not preserved.", "pdf_to_powerpoint": "Creates slides from PDF pages; textual layout is not editable as original.", "pdf_to_excel": "Extracts text rows; table layout may not be preserved."}

@app.get("/api/tools")
def tools_catalog():
    result = []
    for i, t, a, m, handler in TOOL_DEFS:
        fields = TOOL_FIELDS.get(handler, [])
        reason = ""
        if handler == "pdf_to_pdfa" and not ghostscript_binary():
            reason = "Ghostscript belum tersedia."
        if handler in ("word_to_pdf", "powerpoint_to_pdf", "excel_to_pdf") and not office_binary():
            reason = "LibreOffice belum tersedia."
        a = {"word_to_pdf": ".doc,.docx", "powerpoint_to_pdf": ".ppt,.pptx", "excel_to_pdf": ".xls,.xlsx", "jpg_to_pdf": ".jpg,.jpeg,.png", "html_to_pdf": ".html,.htm"}.get(handler, a)
        result.append({"id": i, "title": t, "description": TOOL_DESCRIPTIONS.get(handler, t), "accept": a, "multiple": m, "fields": fields, "available": not bool(reason), "reason": reason, "category": "pdf" if i.startswith("pdf-") else "convert"})
    return jsonify(result)

@app.post("/api/convert/<tool_id>")
def convert_tool(tool_id):
    match = next((x for x in TOOL_DEFS if x[0] == tool_id), None)
    if not match: return fail("Alat tidak ditemukan", 404)
    if match[4] == "pdf_to_pdfa" and not ghostscript_binary():
        return fail("Ghostscript belum tersedia", 503)
    if match[4] in ("word_to_pdf", "powerpoint_to_pdf", "excel_to_pdf") and not office_binary():
        return fail("LibreOffice belum tersedia", 503)
    try:
        return globals()[match[4]]()
    except (ValueError, TypeError, KeyError, IndexError):
        return fail("File atau opsi tidak valid", 400)
    except Exception:
        app.logger.exception("Conversion failed for %s", tool_id)
        return fail("Konversi gagal. Periksa file dan opsi, lalu coba lagi.", 422)
@app.post("/api/pdf-to-pdfa")
def pdf_to_pdfa_legacy():
    return pdf_to_pdfa()


@app.errorhandler(413)
def too_large(_error):
    return fail("Ukuran file terlalu besar (maks 100 MB)", 413)

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5001))
    try:
        from waitress import serve
        serve(app, host="0.0.0.0", port=port)
    except ImportError:
        app.run(host="0.0.0.0", port=port, debug=False)
